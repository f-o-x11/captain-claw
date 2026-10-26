"""docx tool: Word documents built from Markdown and edited as tracked changes.

The redline semantics were checked against Microsoft Word itself (its Accept
All / Reject All give the same text as :func:`docx_ops.resolve`); these tests
pin that behaviour without Word.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx.oxml.ns import qn
from lxml import etree

from captain_claw import docx_ops as D
from captain_claw.tools.docx_tool import DocxTool
from captain_claw.tools.edit import EditTool
from captain_claw.tools.write import WriteTool

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

CONTRACT = """# Services Agreement

This Agreement is made between **Acme Ltd** ("Client") and Beta ("Provider").

## 1. Term

The term is 12 months from the Effective Date. Either party may terminate with 30 days' notice.

1. First obligation
2. Second obligation

Another paragraph.

1. Restarted item
2. And another

| Item | Price |
|------|------:|
| Widget | 10 EUR |
"""


def _doc(tmp_path: Path, md: str = CONTRACT) -> Path:
    doc, _ = D.build_document(md)
    path = tmp_path / "contract.docx"
    doc.save(path)
    return path


def _texts(doc, view: str) -> list[str]:
    return [D.render_paragraph(p, view) for p in D.paragraphs(doc)]


def _para(doc, text: str):
    return next(p for p in D.paragraphs(doc) if text in D.visible_text(p))


def _edit(path: Path, edits, author: str = "Captain"):
    doc = D.open_document(path)
    rv = D.Reviser.for_document(doc, author)
    report = D.apply_edits(doc, edits, rv)
    doc.save(path)
    return D.open_document(path), rv, report


def _resolved(path: Path, accept: bool, **kw) -> list[str]:
    doc = D.open_document(path)
    D.resolve(doc, accept=accept, **kw)
    return [t for t in _texts(doc, "final")]


# ── Building ─────────────────────────────────────────────────────────


def test_markdown_builds_styled_blocks(tmp_path):
    doc = D.open_document(_doc(tmp_path))
    rendered, _ = D.render(doc)
    assert "(H1) Services Agreement" in rendered
    assert "(H2) 1. Term" in rendered
    assert "(List Number) First obligation" in rendered
    assert "(T1 r2c2) 10 EUR" in rendered
    bold = [r for r in _para(doc, "Acme Ltd").iter(qn("w:r")) if D._run_text(r) == "Acme Ltd"][0]
    assert bold.find(f"{qn('w:rPr')}/{qn('w:b')}") is not None


def test_each_numbered_list_restarts(tmp_path):
    doc = D.open_document(_doc(tmp_path))

    def num_id(text):
        return _para(doc, text).find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}").get(qn("w:val"))

    assert num_id("First obligation") == num_id("Second obligation")
    assert num_id("First obligation") != num_id("Restarted item")


def test_inline_markdown_keeps_snake_case_and_arithmetic():
    spans = D.parse_inline("call snake_case_name and 2 * 3 * 4 with *em* and [x](https://a.b)")
    assert "".join(s.text for s in spans) == "call snake_case_name and 2 * 3 * 4 with em and x"
    assert [s.text for s in spans if s.italic] == ["em"]
    assert [s.link for s in spans if s.link] == ["https://a.b"]


# ── Tracked edits ────────────────────────────────────────────────────


def test_replace_marks_only_the_changed_words(tmp_path):
    path = _doc(tmp_path)
    doc, rv, _ = _edit(path, [{"op": "replace", "find": "The term is 12 months",
                               "replace_with": "The term is 24 months"}])
    p = _para(doc, "24 months")
    assert D.render_paragraph(p).startswith("The term is [-12-]{+24+} months")
    dels = list(p.iter(qn("w:del")))
    assert len(dels) == 1 and dels[0].get(qn("w:author")) == "Captain"
    assert dels[0].get(qn("w:date")).endswith("Z")
    assert rv.counts["insertions"] == 1 and rv.counts["deletions"] == 1


def test_a_rewrite_is_one_deletion_and_one_insertion(tmp_path):
    path = _doc(tmp_path)
    doc, rv, _ = _edit(path, [{"op": "replace", "find": "Another paragraph.",
                               "replace_with": "Something else entirely here."}])
    assert D.render_paragraph(_para(doc, "Something")) == \
        "[-Another paragraph.-]{+Something else entirely here.+}"


def test_accept_and_reject_give_final_and_original(tmp_path):
    path = _doc(tmp_path)
    original = [t for t in _texts(D.open_document(path), "final") if t]
    _edit(path, [
        {"op": "format", "find": "Acme Ltd", "italic": True},
        {"op": "replace", "find": "30 days' notice", "replace_with": "60 days’ written notice"},
        {"op": "insert", "after": "Another paragraph.", "text": " Added."},
        {"op": "delete", "find": "Second obligation"},
        {"op": "insert_paragraphs", "after": 3, "content": "New clause.\n\n- a\n- b"},
        {"op": "delete", "paragraph": 2},
        {"op": "set_style", "paragraph": 7, "style": "Heading 2"},
    ])
    edited = D.open_document(path)
    final = [t for t in _texts(edited, "final") if t]
    assert [t for t in _resolved(path, accept=False) if t] == original
    assert [t for t in _texts(edited, "original") if t] == original
    accepted = [t for t in _resolved(path, accept=True) if t]
    assert accepted == final
    assert "New clause." in accepted and not any("This Agreement" in t for t in accepted)
    assert "Another paragraph. Added." in accepted


def test_formatting_change_is_reversible(tmp_path):
    path = _doc(tmp_path)
    _edit(path, [{"op": "format", "find": "Another", "bold": True, "highlight": "yellow"}])
    doc = D.open_document(path)
    change = next(doc.element.body.iter(qn("w:rPrChange")))
    assert change.get(qn("w:author")) == "Captain"
    rejected = D.open_document(path)
    D.resolve(rejected, accept=False)
    run = [r for r in _para(rejected, "Another").iter(qn("w:r")) if "Another" in D._run_text(r)][0]
    assert run.find(f"{qn('w:rPr')}/{qn('w:b')}") is None
    assert run.find(f"{qn('w:rPr')}/{qn('w:highlight')}") is None


def test_style_change_rejects_back(tmp_path):
    path = _doc(tmp_path)
    _edit(path, [{"op": "set_style", "paragraph": 6, "style": "Heading 2"}])
    doc = D.open_document(path)
    assert D.list_changes(doc)[0].kind == "paragraph style"
    D.resolve(doc, accept=False)
    ppr = D.paragraphs(doc)[5].find(qn("w:pPr"))
    style = ppr.find(qn("w:pStyle")) if ppr is not None else None
    assert style is None or style.get(qn("w:val")) != "Heading2"


def test_deleting_your_own_insertion_leaves_no_trace(tmp_path):
    path = _doc(tmp_path)
    _edit(path, [{"op": "insert", "after": "Another paragraph.", "text": " Temporary."}])
    doc, _, _ = _edit(path, [{"op": "delete", "find": " Temporary."}])
    p = _para(doc, "Another paragraph.")
    assert D.render_paragraph(p) == "Another paragraph."
    assert not list(p.iter(qn("w:del")))


def _foreign_insertion(path: Path, text: str, after: str) -> None:
    """Simulate another reviewer's tracked insertion (as Word writes it)."""
    doc = D.open_document(path)
    p = _para(doc, after)
    ins = etree.SubElement(p, qn("w:ins"))
    ins.set(qn("w:id"), "900")
    ins.set(qn("w:author"), "Ana Reviewer")
    ins.set(qn("w:date"), "2026-01-01T00:00:00Z")
    r = etree.SubElement(ins, qn("w:r"))
    t = etree.SubElement(r, qn("w:t"))
    t.text = text
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    doc.save(path)


def test_editing_inside_another_authors_insertion(tmp_path):
    path = _doc(tmp_path)
    _foreign_insertion(path, " Late payments accrue interest.", "Another paragraph.")
    doc, _, _ = _edit(path, [{"op": "replace", "find": "Late payments",
                              "replace_with": "Overdue amounts"}])
    p = _para(doc, "Overdue amounts")
    # Our deletion sits inside Ana's insertion; our insertion is never nested in hers.
    our_del = [d for d in p.iter(qn("w:del")) if d.get(qn("w:author")) == "Captain"][0]
    assert our_del.getparent().tag == qn("w:ins")
    for ins in p.iter(qn("w:ins")):
        if ins.get(qn("w:author")) == "Captain":
            assert ins.getparent().tag != qn("w:ins")
    by = {c.author for c in D.list_changes(doc)}
    assert by == {"Ana Reviewer", "Captain"}
    # Accepting only Captain's changes keeps Ana's for her own review.
    D.resolve(doc, accept=True, author="captain")
    left = D.list_changes(doc)
    assert left and {c.author for c in left} == {"Ana Reviewer"}
    assert "Overdue amounts accrue interest." in D.visible_text(_para(doc, "Overdue"))


def test_accept_by_id(tmp_path):
    path = _doc(tmp_path)
    doc, _, _ = _edit(path, [{"op": "replace", "find": "12 months", "replace_with": "24 months"}])
    changes = D.list_changes(doc)
    D.resolve(doc, accept=True, ids=[changes[0].id])
    left = {c.id for c in D.list_changes(doc)}
    assert changes[0].id not in left and len(left) == len(changes) - 1


def test_comment_is_a_real_word_comment(tmp_path):
    path = _doc(tmp_path)
    doc, rv, _ = _edit(path, [{"op": "comment", "find": "Acme Ltd", "text": "Confirm the legal name."}])
    [c] = D.list_comments(doc)
    assert (c.author, c.text, c.anchor) == ("Captain", "Confirm the legal name.", "Acme Ltd")
    assert rv.counts["comments"] == 1


def test_inserted_paragraph_after_a_heading_takes_the_next_style(tmp_path):
    path = _doc(tmp_path)
    doc, _, _ = _edit(path, [{"op": "insert_paragraphs", "after": 3, "content": "Body text."}])
    p = _para(doc, "Body text.")
    style = p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
    assert style is None or not style.get(qn("w:val")).startswith("Heading")
    assert D._mark_state(p)[0] == "ins"


def test_find_folds_typographic_quotes(tmp_path):
    path = _doc(tmp_path, "He said “yes” and it’s done.")
    doc, _, _ = _edit(path, [{"op": "replace", "find": "it's done", "replace_with": "it is done"}])
    assert "it is done" in D.visible_text(D.paragraphs(doc)[0])


@pytest.mark.parametrize("edit, message", [
    ({"op": "replace", "find": "obligation", "replace_with": "x"}, "occurs 2 times"),
    ({"op": "replace", "find": "no such words", "replace_with": "x"}, "Text not found"),
    ({"op": "replace", "find": "12 months"}, "replace_with"),
    ({"op": "explode"}, "unknown op"),
    ({"op": "delete", "paragraph": 999}, "doesn't exist"),
    ({"op": "set_style", "paragraph": 1, "style": "No Such Style"}, "no paragraph style"),
])
def test_bad_edits_explain_themselves(tmp_path, edit, message):
    path = _doc(tmp_path)
    before = path.read_bytes()
    doc = D.open_document(path)
    with pytest.raises(D.DocxError, match=message):
        D.apply_edits(doc, [edit], D.Reviser.for_document(doc, "Captain"))
    assert path.read_bytes() == before


def test_field_crossing_is_refused(tmp_path):
    doc, _ = D.build_document("Page X here.")
    p = D.paragraphs(doc)[0]
    run = next(p.iter(qn("w:r")))
    field_result = D._split_run(run, 5)          # "Page " | "X here."
    D._split_run(field_result, 1)                # "X" | " here."

    def fld(kind):
        r = etree.Element(qn("w:r"))
        etree.SubElement(r, qn("w:fldChar")).set(qn("w:fldCharType"), kind)
        return r

    field_result.addprevious(fld("begin"))
    field_result.addnext(fld("end"))
    rv = D.Reviser.for_document(doc, "Captain")
    with pytest.raises(D.DocxError, match="field"):
        D.apply_edits(doc, [{"op": "replace", "find": "Page X", "replace_with": "Sheet Y"}], rv)
    D.apply_edits(doc, [{"op": "replace", "find": "here", "replace_with": "there"}], rv)
    assert D.visible_text(p) == "Page X there."


def test_track_revisions_lands_in_schema_order(tmp_path):
    doc = D.open_document(_doc(tmp_path))
    D.set_tracking(doc, True)
    names = [etree.QName(ch).localname for ch in doc.settings.element]
    assert "trackRevisions" in names
    if "defaultTabStop" in names:
        assert names.index("trackRevisions") < names.index("defaultTabStop")
    assert D.tracking_on(doc)


# ── The tool ─────────────────────────────────────────────────────────


@pytest.fixture
def workspace(tmp_path):
    (tmp_path / "saved" / "downloads" / "s1").mkdir(parents=True)
    kw = {"_runtime_base_path": tmp_path, "_saved_base_path": tmp_path / "saved", "_session_id": "s1"}
    return tmp_path, kw


async def test_an_upload_is_never_changed(workspace):
    root, kw = workspace
    upload = root / "saved" / "downloads" / "s1" / "Contract.docx"
    upload.write_bytes(_doc(root).read_bytes())
    before = upload.read_bytes()
    tool = DocxTool()
    result = await tool.execute(action="edit", path="saved/downloads/s1/Contract.docx",
                                edits=[{"op": "replace", "find": "12 months", "replace_with": "24 months"}],
                                author="Captain", **kw)
    assert result.success, result.error
    assert upload.read_bytes() == before
    copy_path = root / "saved" / "output" / "s1" / "Contract.docx"
    assert copy_path.is_file() and str(copy_path) in result.content
    assert "[-12-]{+24+}" in result.content
    assert D.tracking_on(D.open_document(copy_path))

    again = await tool.execute(action="edit", path="saved/downloads/s1/Contract.docx",
                               edits=[{"op": "replace", "find": "Another", "replace_with": "One more"}], **kw)
    assert not again.success and "already exists" in again.error

    in_place = await tool.execute(action="edit", path="saved/output/s1/Contract.docx",
                                  edits=[{"op": "delete", "find": "Another "}], **kw)
    assert in_place.success and "Saved in place" in in_place.content and "Backup:" in in_place.content


async def test_create_read_and_resolve_through_the_tool(workspace):
    root, kw = workspace
    tool = DocxTool()
    made = await tool.execute(action="create", path="Offer.docx", content="# Offer\n\nPrice is 10 EUR.",
                              title="Offer", page_numbers=True, **kw)
    assert made.success, made.error
    target = root / "saved" / "tmp" / "s1" / "Offer.docx"
    assert target.is_file()
    dup = await tool.execute(action="create", path="Offer.docx", content="# x", **kw)
    assert not dup.success and "already exists" in dup.error

    read = await tool.execute(action="read", path="Offer.docx", **kw)
    assert read.success and "[1] (H1) Offer" in read.content and "Footer: Page" in read.content

    await tool.execute(action="edit", path="Offer.docx", author="Captain",
                       edits='[{"op": "replace", "find": "10 EUR", "replace_with": "12 EUR"}]', **kw)
    changes = await tool.execute(action="changes", path="Offer.docx", **kw)
    assert "deletion by Captain" in changes.content and "'10'" in changes.content
    need = await tool.execute(action="accept", path="Offer.docx", **kw)
    assert not need.success and "all=true" in need.error
    done = await tool.execute(action="accept", path="Offer.docx", all=True, **kw)
    assert done.success and "0 left" in done.content
    assert "12 EUR" in D.visible_text(D.paragraphs(D.open_document(target))[1])


async def test_write_and_edit_refuse_word_files(workspace):
    root, kw = workspace
    wrote = await WriteTool().execute(path="report.docx", content="hello", **kw)
    assert not wrote.success and "docx tool" in wrote.content
    edited = await EditTool().execute(path="x.docx", action="replace_string",
                                      old_string="a", new_string="b", **kw)
    assert not edited.success and "docx tool" in edited.error
