"""Word (.docx) documents for agents, without scripts.

Three jobs, all done on the document's own OOXML so a reviewer in Word sees
exactly what a careful human editor would have produced:

* build a document from Markdown (:func:`build_document`): headings, lists
  whose numbering restarts per list, tables, quotes, code, pictures with
  captions, links, a table of contents and page numbers;
* show a document paragraph by paragraph (:func:`render`) with its tracked
  changes and comments inline, so edits can be aimed at numbered paragraphs;
* change it the way Word does with Track Changes on (:func:`apply_edits`):
  every insertion, deletion, formatting and style change is a revision with an
  author and a date (``w:ins``, ``w:del``, ``w:rPrChange``, ``w:pPrChange``),
  a replacement is cut down to the words that actually changed, and comments
  are real Word comments. :func:`resolve` accepts or rejects revisions the way
  Word's Accept / Reject buttons do.

Nothing here decides where files live; ``tools/docx_tool.py`` owns paths,
backups and saving.
"""

from __future__ import annotations

import copy
import difflib
import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Callable, Iterable

from docx.document import Document as DocxDocument
from docx.enum.text import WD_BREAK
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree

# ── Tags ─────────────────────────────────────────────────────────────

P, R, T, DEL_TEXT = qn("w:p"), qn("w:r"), qn("w:t"), qn("w:delText")
INSTR, DEL_INSTR = qn("w:instrText"), qn("w:delInstrText")
PPR, RPR, TBL, TR, TC = qn("w:pPr"), qn("w:rPr"), qn("w:tbl"), qn("w:tr"), qn("w:tc")
TRPR, TCPR = qn("w:trPr"), qn("w:tcPr")
INS, DEL = qn("w:ins"), qn("w:del")
MOVE_FROM, MOVE_TO = qn("w:moveFrom"), qn("w:moveTo")
RPR_CHANGE, PPR_CHANGE = qn("w:rPrChange"), qn("w:pPrChange")
TAB, PTAB, BR, CR = qn("w:tab"), qn("w:ptab"), qn("w:br"), qn("w:cr")
NB_HYPHEN = qn("w:noBreakHyphen")
FLD_CHAR = qn("w:fldChar")
SECT_PR = qn("w:sectPr")
TXBX = qn("w:txbxContent")
CRS, CRE = qn("w:commentRangeStart"), qn("w:commentRangeEnd")
COMMENT_REF = qn("w:commentReference")
CELL_INS, CELL_DEL = qn("w:cellIns"), qn("w:cellDel")
W_ID, W_AUTHOR, W_DATE, W_VAL = qn("w:id"), qn("w:author"), qn("w:date"), qn("w:val")
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

TRACK_CONTAINERS = frozenset({INS, DEL, MOVE_FROM, MOVE_TO})
HIDDEN_CONTAINERS = frozenset({DEL, MOVE_FROM})
SHOWN_CONTAINERS = frozenset({INS, MOVE_TO})
# Property-change revisions: the element, and the children of its parent a
# reject keeps — those that come before the properties, and those after. The
# rest is replaced by the old properties the change element holds.
_NONE: frozenset[str] = frozenset()
PROPERTY_CHANGES: dict[str, tuple[frozenset[str], frozenset[str]]] = {
    RPR_CHANGE: (frozenset({INS, DEL, MOVE_FROM, MOVE_TO}), _NONE),
    PPR_CHANGE: (_NONE, frozenset({RPR, SECT_PR})),
    qn("w:tblPrChange"): (_NONE, _NONE),
    qn("w:trPrChange"): (_NONE, frozenset({INS, DEL})),
    qn("w:tcPrChange"): (_NONE, frozenset({CELL_INS, CELL_DEL, qn("w:cellMerge")})),
    qn("w:sectPrChange"): (frozenset({qn("w:headerReference"), qn("w:footerReference")}), _NONE),
    qn("w:tblGridChange"): (_NONE, _NONE),
}

DOCX_TYPES = {
    CT.WML_DOCUMENT_MAIN,
    "application/vnd.ms-word.document.macroEnabled.main+xml",
}
TEMPLATE_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
    "application/vnd.ms-word.template.macroEnabledTemplate.main+xml",
}

# Characters the model types for what Word stores as typographic ones.
_FOLD = str.maketrans({
    "‘": "'", "’": "'", "‚": "'", "‛": "'", "′": "'",
    "“": '"', "”": '"', "„": '"', "″": '"',
    " ": " ", " ": " ", " ": " ", " ": " ",
    "‐": "-", "‑": "-", "‒": "-", "–": "-", "—": "-",
})


class DocxError(ValueError):
    """A request the document can't satisfy; the message is for the model."""


# ── Opening ──────────────────────────────────────────────────────────


def open_document(path: str | Path, *, template: bool = False) -> DocxDocument:
    """Open a .docx (or .docm). With *template*, a .dotx/.dotm is accepted
    too and turned into an ordinary document."""
    from docx.package import Package

    try:
        part = Package.open(str(path)).main_document_part
    except Exception as exc:  # zip / xml damage
        raise DocxError(f"Not a readable Word document ({type(exc).__name__}: {exc}).") from exc
    content_type = part.content_type
    if content_type in TEMPLATE_TYPES and template:
        part._content_type = CT.WML_DOCUMENT_MAIN
    elif content_type not in DOCX_TYPES:
        raise DocxError(
            "Not a Word document (.docx). Legacy .doc and other formats must be "
            "converted to .docx first.")
    return part.document


# ── Small XML helpers ────────────────────────────────────────────────


def _el(name: str) -> Any:
    return OxmlElement(name)


def _text_el(text: str, tag: str = "w:t") -> Any:
    t = _el(tag)
    t.text = text
    t.set(XML_SPACE, "preserve")
    return t


def _has_ancestor(el: Any, tags: frozenset[str] | set[str], stop: Any) -> bool:
    node = el.getparent()
    while node is not None and node is not stop:
        if node.tag in tags:
            return True
        node = node.getparent()
    return False


def _attached(el: Any, root: Any) -> bool:
    node = el
    while node is not None:
        if node is root:
            return True
        node = node.getparent()
    return False


def _now_iso() -> str:
    return datetime.now(UTC).replace(microsecond=0).strftime("%Y-%m-%dT%H:%M:%SZ")


def initials_of(author: str) -> str:
    words = [w for w in re.split(r"[\s._-]+", author or "") if w]
    return "".join(w[0] for w in words[:3]).upper() or "AI"


@dataclass
class Reviser:
    """Stamps revisions with one author and date, and counts them."""

    author: str
    date: str = field(default_factory=_now_iso)
    next_id: int = 1
    counts: dict[str, int] = field(default_factory=lambda: {
        "insertions": 0, "deletions": 0, "formatting": 0, "styles": 0,
        "paragraphs_inserted": 0, "paragraphs_deleted": 0, "comments": 0,
    })

    @classmethod
    def for_document(cls, doc: DocxDocument, author: str) -> "Reviser":
        top = 0
        for el in doc.element.iter():
            raw = el.get(W_ID)
            if raw is not None and raw.lstrip("-").isdigit():
                top = max(top, int(raw))
        return cls(author=author, next_id=top + 1)

    def new_id(self) -> str:
        value = self.next_id
        self.next_id += 1
        return str(value)

    def stamp(self, el: Any) -> Any:
        el.set(W_ID, self.new_id())
        el.set(W_AUTHOR, self.author)
        el.set(W_DATE, self.date)
        return el


# ── Paragraphs, runs and visible text ────────────────────────────────


def paragraphs(doc: DocxDocument) -> list[Any]:
    """Every body paragraph (tables included) in document order, numbered
    1.. by position in this list. Text-box paragraphs are left out — their
    text is drawn inside another paragraph's run."""
    body = doc.element.body
    return [p for p in body.iter(P) if not _has_ancestor(p, {TXBX}, body)]


def _own_runs(p: Any) -> list[Any]:
    """Runs of *p* itself, in order (not those of a nested text-box paragraph)."""
    out = []
    for r in p.iter(R):
        node, ok = r.getparent(), True
        while node is not None and node is not p:
            if node.tag in (P, TXBX):
                ok = False
                break
            node = node.getparent()
        if ok:
            out.append(r)
    return out


def _run_state(r: Any, p: Any) -> str:
    """"del", "ins" or "" for a run of paragraph *p*."""
    node, state = r.getparent(), ""
    while node is not None and node is not p:
        if node.tag in HIDDEN_CONTAINERS:
            return "del"
        if node.tag in SHOWN_CONTAINERS:
            state = "ins"
        node = node.getparent()
    return state


def _pieces(r: Any) -> list[tuple[Any, str]]:
    """The run's content children with the text each contributes."""
    out: list[tuple[Any, str]] = []
    for ch in r:
        tag = ch.tag
        if tag == RPR:
            continue
        if tag in (T, DEL_TEXT):
            out.append((ch, ch.text or ""))
        elif tag in (TAB, PTAB):
            out.append((ch, "\t"))
        elif tag in (BR, CR):
            out.append((ch, "\n"))
        elif tag == NB_HYPHEN:
            out.append((ch, "-"))
        else:  # fldChar, instrText, drawing, sym, softHyphen, footnote refs …
            out.append((ch, ""))
    return out


def _run_text(r: Any) -> str:
    return "".join(text for _, text in _pieces(r))


def _segments(p: Any) -> list[tuple[Any, int, int]]:
    """Visible runs of *p* with their [start, end) in the paragraph's text."""
    segs, pos = [], 0
    for r in _own_runs(p):
        if _run_state(r, p) == "del":
            continue
        n = len(_run_text(r))
        segs.append((r, pos, pos + n))
        pos += n
    return segs


def visible_text(p: Any) -> str:
    return "".join(_run_text(r) for r, _, _ in _segments(p))


def _split_run(r: Any, offset: int) -> Any | None:
    """Split run *r* at *offset* of its text; the right half (a new run with
    the same properties) is placed right after it and returned."""
    pieces = _pieces(r)
    total = sum(len(t) for _, t in pieces)
    if offset <= 0 or offset >= total:
        return None
    right = _el("w:r")
    for key, value in r.attrib.items():
        right.set(key, value)
    rpr = r.find(RPR)
    if rpr is not None:
        right.append(copy.deepcopy(rpr))
    pos, moving = 0, False
    for ch, text in pieces:
        n = len(text)
        if moving:
            right.append(ch)
            continue
        if pos + n <= offset:
            pos += n
            moving = pos == offset and n > 0
            continue
        cut = offset - pos  # 0 < cut < n: only text elements are this long
        tag = "w:delText" if ch.tag == DEL_TEXT else "w:t"
        ch.text = text[:cut]
        ch.set(XML_SPACE, "preserve")
        right.append(_text_el(text[cut:], tag))
        moving = True
    r.addnext(right)
    return right


def _isolate(p: Any, start: int, end: int) -> list[Any]:
    """Split runs so [start, end) of *p*'s visible text is whole runs; return them."""
    for edge in (end, start):
        for r, s, e in _segments(p):
            if s < edge < e:
                _split_run(r, edge - s)
                break
    runs = [r for r, s, e in _segments(p)
            if s >= start and e <= end and (e > s or start < s < end)]
    depth = 0
    for r in runs:
        for fc in r.iter(FLD_CHAR):
            kind = fc.get(qn("w:fldCharType"))
            depth += 1 if kind == "begin" else -1 if kind == "end" else 0
            if depth < 0:
                break
    if depth != 0:
        raise DocxError(
            "That text runs into a field (page number, date, cross-reference or "
            "similar). Change the text before or after the field instead.")
    return runs


def _sibling_groups(runs: list[Any]) -> list[list[Any]]:
    groups: list[list[Any]] = []
    for r in runs:
        if groups and groups[-1][-1].getnext() is r:
            groups[-1].append(r)
        else:
            groups.append([r])
    return groups


# ── Revision primitives ──────────────────────────────────────────────


def _to_deleted(r: Any) -> None:
    for t in list(r.iter(T)):
        t.tag = DEL_TEXT
    for it in list(r.iter(INSTR)):
        it.tag = DEL_INSTR


def _to_live(el: Any) -> None:
    for t in list(el.iter(DEL_TEXT)):
        t.tag = T
    for it in list(el.iter(DEL_INSTR)):
        it.tag = INSTR


def _track_delete(runs: list[Any], rv: Reviser) -> None:
    """Mark *runs* deleted. Text this author inserted earlier simply goes —
    Word does the same when you delete your own tracked insertion."""
    for group in _sibling_groups(runs):
        parent = group[0].getparent()
        if parent.tag == INS and parent.get(W_AUTHOR) == rv.author:
            for r in group:
                parent.remove(r)
            if not len(parent):
                parent.getparent().remove(parent)
            continue
        wrapper = rv.stamp(_el("w:del"))
        group[0].addprevious(wrapper)
        for r in group:
            wrapper.append(r)
            _to_deleted(r)
        rv.counts["deletions"] += 1


def _content_runs(text: str, rpr: Any | None) -> list[Any]:
    """Runs carrying *text*, tabs and line breaks included, formatted like *rpr*."""
    r = _el("w:r")
    if rpr is not None:
        props = copy.deepcopy(rpr)
        for change in props.findall(RPR_CHANGE):
            props.remove(change)
        r.append(props)
    for i, line in enumerate(text.split("\n")):
        if i:
            r.append(_el("w:br"))
        for j, chunk in enumerate(line.split("\t")):
            if j:
                r.append(_el("w:tab"))
            if chunk:
                r.append(_text_el(chunk))
    return [r]


_MARKER_TAG = "{urn:captain-claw:docx}marker"


def _lift_marker(marker: Any, rv: Reviser) -> None:
    """Move *marker* out of any tracked-change container, splitting it."""
    parent = marker.getparent()
    while parent is not None and parent.tag in TRACK_CONTAINERS:
        following = list(marker.itersiblings())
        parent.remove(marker)
        if following:
            clone = parent.makeelement(parent.tag, dict(parent.attrib))
            clone.set(W_ID, rv.new_id())
            for el in following:
                clone.append(el)
            parent.addnext(clone)
        parent.addnext(marker)
        if not len(parent):
            parent.getparent().remove(parent)
        parent = marker.getparent()


def _insert_at_marker(marker: Any, text: str, rpr: Any | None, rv: Reviser) -> None:
    _lift_marker(marker, rv)
    if text:
        ins = rv.stamp(_el("w:ins"))
        for r in _content_runs(text, rpr):
            ins.append(r)
        marker.addprevious(ins)
        rv.counts["insertions"] += 1
    marker.getparent().remove(marker)


def _marker_at(p: Any, pos: int) -> tuple[Any, Any | None]:
    """A marker at visible offset *pos* of *p*, and the run whose formatting
    text typed there would take."""
    marker = etree.Element(_MARKER_TAG)
    segs = _segments(p)
    if pos > 0:
        for r, s, e in segs:
            if s < pos < e:
                _split_run(r, pos - s)
                break
        segs = _segments(p)
        before = [r for r, s, e in segs if e == pos and e > s]
        if before:
            before[-1].addnext(marker)
            return marker, before[-1]
    after = [r for r, s, e in segs if s >= pos]
    if after:
        after[0].addprevious(marker)
        return marker, after[0]
    if segs:
        segs[-1][0].addnext(marker)
        return marker, segs[-1][0]
    ppr = p.find(PPR)
    if ppr is not None:
        ppr.addnext(marker)
    else:
        p.insert(0, marker)
    return marker, None


def _rpr_of(r: Any | None) -> Any | None:
    return None if r is None else r.find(RPR)


_TOKEN_RE = re.compile(r"\w+|\s+|[^\w\s]", re.UNICODE)


def _diff_ops(old: str, new: str) -> list[tuple[str, int, int, str]]:
    """Word-level changes turning *old* into *new*: (kind, start, end, text)
    with offsets into *old*. A near-total rewrite comes back as one change —
    a clean redline reads better than confetti."""
    a, b = _TOKEN_RE.findall(old), _TOKEN_RE.findall(new)
    matcher = difflib.SequenceMatcher(None, a, b, autojunk=False)
    if matcher.ratio() < 0.5:
        return [("replace", 0, len(old), new)]
    offsets = [0]
    for tok in a:
        offsets.append(offsets[-1] + len(tok))
    out: list[tuple[str, int, int, str]] = []
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        out.append((tag, offsets[i1], offsets[i2], "".join(b[j1:j2])))
        # Two changes with only spaces between them read as one change.
        while (len(out) >= 3 and out[-1][0] != "equal" and out[-3][0] != "equal"
               and out[-2][0] == "equal" and not old[out[-2][1]:out[-2][2]].strip()):
            first, gap, last = out[-3], out[-2], out[-1]
            out[-3:] = [("replace", first[1], last[2],
                         first[3] + old[gap[1]:gap[2]] + last[3])]
    return [op for op in out if op[0] != "equal"]


def _track_replace(p: Any, start: int, end: int, new: str, rv: Reviser) -> None:
    old = visible_text(p)[start:end]
    changes = _diff_ops(old, new) if new else [("delete", 0, len(old), "")]
    for kind, s, e, text in reversed(changes):
        a, b = start + s, start + e
        if kind == "insert":
            marker, ref = _marker_at(p, a)
            _insert_at_marker(marker, text, _rpr_of(ref), rv)
            continue
        runs = _isolate(p, a, b)
        if not runs:
            continue
        rpr = _rpr_of(runs[0])
        marker = etree.Element(_MARKER_TAG)
        runs[-1].addnext(marker)
        _track_delete(runs, rv)
        _insert_at_marker(marker, text if kind == "replace" else "", rpr, rv)


# ── Paragraph marks ──────────────────────────────────────────────────

# w:pPr children in schema order — a new w:rPr goes before what follows it.
_PPR_AFTER_RPR = (SECT_PR, PPR_CHANGE)


def _mark_rpr(p: Any) -> Any:
    ppr = p.get_or_add_pPr()
    rpr = ppr.find(RPR)
    if rpr is None:
        rpr = _el("w:rPr")
        for tag in _PPR_AFTER_RPR:
            nxt = ppr.find(tag)
            if nxt is not None:
                nxt.addprevious(rpr)
                break
        else:
            ppr.append(rpr)
    return rpr


def _mark_paragraph(p: Any, kind: str, rv: Reviser) -> None:
    """Flag *p*'s paragraph mark as inserted or deleted ("ins" / "del")."""
    rpr = _mark_rpr(p)
    for old in rpr.findall(INS) + rpr.findall(DEL):
        rpr.remove(old)
    marker = rv.stamp(_el(f"w:{kind}"))
    rpr.insert(0, marker)


def _mark_state(p: Any) -> tuple[str, Any | None]:
    ppr = p.find(PPR)
    rpr = ppr.find(RPR) if ppr is not None else None
    if rpr is None:
        return "", None
    for tag, kind in ((INS, "ins"), (DEL, "del"), (MOVE_TO, "ins"), (MOVE_FROM, "del")):
        el = rpr.find(tag)
        if el is not None:
            return kind, el
    return "", None


def _wrap_new_runs(container: Any, rv: Reviser) -> None:
    """Wrap every run under *container* in w:ins (consecutive siblings together)."""
    runs = [r for r in container.iter(R)
            if not _has_ancestor(r, set(TRACK_CONTAINERS), container)]
    for group in _sibling_groups(runs):
        ins = rv.stamp(_el("w:ins"))
        group[0].addprevious(ins)
        for r in group:
            ins.append(r)


def _track_new_block(block: Any, rv: Reviser) -> None:
    """Mark a freshly built paragraph or table as a tracked insertion."""
    paras = [block] if block.tag == P else list(block.iter(P))
    for p in paras:
        _wrap_new_runs(p, rv)
        _mark_paragraph(p, "ins", rv)
    if block.tag == TBL:
        for tr in block.iter(TR):
            trpr = tr.find(TRPR)
            if trpr is None:
                trpr = _el("w:trPr")
                tblpr_ex = tr.find(qn("w:tblPrEx"))
                if tblpr_ex is not None:
                    tblpr_ex.addnext(trpr)
                else:
                    tr.insert(0, trpr)
            trpr.append(rv.stamp(_el("w:ins")))


def _next_paragraph(p: Any) -> Any | None:
    """The paragraph right after *p* in the same container (bookmarks and
    other markers between them skipped), or None at a table / the end."""
    node = p.getnext()
    while node is not None:
        if node.tag == P:
            return node
        if node.tag in (TBL, SECT_PR, qn("w:sdt")):
            return None
        node = node.getnext()
    return None


def _has_content(p: Any) -> bool:
    return any(ch.tag != RPR for r in _own_runs(p) for ch in r)


def _drop_mark(p: Any) -> None:
    """Remove *p*'s paragraph mark: its content joins the next paragraph,
    which keeps its own formatting (Word's behaviour). With no paragraph to
    join (a table or the end follows), an emptied paragraph just goes."""
    nxt = _next_paragraph(p)
    content = [ch for ch in p if ch.tag != PPR]
    if nxt is None:
        prev = p.getprevious()
        if not _has_content(p) and prev is not None and prev.tag == P:
            p.getparent().remove(p)
        return
    nppr = nxt.find(PPR)
    anchor_index = 0 if nppr is None else list(nxt).index(nppr) + 1
    for offset, ch in enumerate(content):
        nxt.insert(anchor_index + offset, ch)
    p.getparent().remove(p)


# ── Formatting and style revisions ───────────────────────────────────

HIGHLIGHTS = {
    "yellow", "green", "cyan", "magenta", "blue", "red", "darkBlue", "darkCyan",
    "darkGreen", "darkMagenta", "darkRed", "darkYellow", "darkGray", "lightGray",
    "black", "white",
}


def _apply_font(r: Any, props: dict[str, Any]) -> None:
    font = Run(r, None).font
    for key in ("bold", "italic", "strike"):
        if key in props:
            setattr(font, key, bool(props[key]))
    if "underline" in props:
        font.underline = bool(props["underline"])
    if "size" in props and props["size"]:
        font.size = Pt(float(props["size"]))
    if "color" in props and props["color"]:
        font.color.rgb = RGBColor.from_string(str(props["color"]).lstrip("#").upper())
    if "highlight" in props:
        rpr = r.get_or_add_rPr()
        for old in rpr.findall(qn("w:highlight")):
            rpr.remove(old)
        value = props["highlight"]
        if value:
            name = {h.lower(): h for h in HIGHLIGHTS}.get(str(value).lower())
            if name is None:
                raise DocxError(f"highlight must be one of: {', '.join(sorted(HIGHLIGHTS))}")
            el = _el("w:highlight")
            el.set(W_VAL, name)
            # w:highlight sits after w:sz/w:szCs and before w:u in w:rPr.
            after = [qn(f"w:{n}") for n in (
                "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em",
                "lang", "eastAsianLayout", "specVanish", "oMath")]
            for tag in after:
                nxt = rpr.find(tag)
                if nxt is not None:
                    nxt.addprevious(el)
                    break
            else:
                rpr.append(el)


def _track_format(runs: list[Any], props: dict[str, Any], rv: Reviser) -> None:
    for r in runs:
        rpr = r.get_or_add_rPr()
        baseline = rpr.find(RPR_CHANGE)
        if baseline is not None:
            rpr.remove(baseline)
        before = copy.deepcopy(rpr)
        _apply_font(r, props)
        rpr = r.get_or_add_rPr()
        if baseline is None:
            baseline = rv.stamp(_el("w:rPrChange"))
            old = _el("w:rPr")
            for ch in list(before):
                old.append(ch)
            baseline.append(old)
        rpr.append(baseline)
    if runs:
        rv.counts["formatting"] += 1


def _track_style(p: Any, style: Any, rv: Reviser) -> None:
    ppr = p.get_or_add_pPr()
    baseline = ppr.find(PPR_CHANGE)
    if baseline is not None:
        ppr.remove(baseline)
    else:
        baseline = rv.stamp(_el("w:pPrChange"))
        old = _el("w:pPr")
        for ch in ppr:
            if ch.tag not in (RPR, SECT_PR):
                old.append(copy.deepcopy(ch))
        baseline.append(old)
    ppr.style = style.style_id
    ppr.append(baseline)
    rv.counts["styles"] += 1


# ── Finding text ─────────────────────────────────────────────────────


@dataclass
class Hit:
    p: Any
    start: int
    end: int


def find_text(paras: list[Any], needle: str, scope: list[Any] | None = None) -> list[Hit]:
    """Where *needle* occurs in the visible text, exactly first, then with
    typographic quotes, dashes and spaces folded to plain ones."""
    if not needle:
        return []
    pool = scope if scope is not None else paras
    hits = _scan(pool, needle, fold=False)
    return hits if hits else _scan(pool, needle, fold=True)


def _scan(pool: list[Any], needle: str, *, fold: bool) -> list[Hit]:
    target = needle.translate(_FOLD) if fold else needle
    hits = []
    for p in pool:
        text = visible_text(p)
        hay = text.translate(_FOLD) if fold else text
        start = hay.find(target)
        while start != -1:
            hits.append(Hit(p, start, start + len(target)))
            start = hay.find(target, start + max(1, len(target)))
    return hits


def _not_found(paras: list[Any], needle: str) -> str:
    if "\n" in needle.strip():
        return ("`find` must sit inside one paragraph; it contains a line break. "
                "Use one edit per paragraph, or delete whole paragraphs by number.")
    words = [w for w in re.findall(r"\w{4,}", needle)][:3]
    near = []
    for i, p in enumerate(paras, 1):
        text = visible_text(p)
        if words and all(w.lower() in text.lower() for w in words):
            near.append(i)
    hint = (f" Paragraphs with similar words: {', '.join(f'[{n}]' for n in near[:8])} — "
            "read them and copy the text exactly.") if near else (
            " Read the document and copy the text exactly as it appears.")
    return f"Text not found: {needle[:120]!r}.{hint}"


# ── Inline Markdown ──────────────────────────────────────────────────


@dataclass
class Span:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False
    strike: bool = False
    link: str = ""
    brk: bool = False


_INLINE_RE = re.compile(
    r"(?P<esc>\\[\\`*_{}\[\]()#+\-.!~<>|])"
    r"|(?P<code>`+)(?P<code_text>.+?)(?P=code)"
    r"|\[(?P<link_text>[^\]]+)\]\((?P<link_url>[^)\s]+)(?:\s+\"[^\"]*\")?\)"
    r"|(?P<bi>\*\*\*)(?=\S)(?P<bi_text>.+?)(?<=\S)\*\*\*"
    r"|(?P<b>\*\*|__)(?=\S)(?P<b_text>.+?)(?<=\S)(?P=b)"
    r"|\*(?=\S)(?P<i_text>.+?)(?<=\S)\*"
    r"|(?<![\w])_(?=\S)(?P<u_text>.+?)(?<=\S)_(?![\w])"
    r"|~~(?=\S)(?P<s_text>.+?)(?<=\S)~~"
    r"|(?P<br><br\s*/?>)",
    re.DOTALL,
)


def parse_inline(text: str, base: Span | None = None) -> list[Span]:
    base = base or Span("")
    spans: list[Span] = []

    def plain(chunk: str) -> None:
        if chunk:
            spans.append(Span(chunk, base.bold, base.italic, base.code, base.strike, base.link))

    pos = 0
    for m in _INLINE_RE.finditer(text):
        plain(text[pos:m.start()])
        pos = m.end()
        g = m.groupdict()
        if g["esc"]:
            plain(g["esc"][1:])
        elif g["code"]:
            spans.append(Span(g["code_text"], base.bold, base.italic, True, base.strike, base.link))
        elif g["link_text"] is not None:
            inner = Span("", base.bold, base.italic, base.code, base.strike, g["link_url"])
            spans.extend(parse_inline(g["link_text"], inner))
        elif g["bi"]:
            spans.extend(parse_inline(g["bi_text"], Span(
                "", True, True, base.code, base.strike, base.link)))
        elif g["b"]:
            spans.extend(parse_inline(g["b_text"], Span(
                "", True, base.italic, base.code, base.strike, base.link)))
        elif g["i_text"] is not None or g["u_text"] is not None:
            inner_text = g["i_text"] if g["i_text"] is not None else g["u_text"]
            spans.extend(parse_inline(inner_text, Span(
                "", base.bold, True, base.code, base.strike, base.link)))
        elif g["s_text"] is not None:
            spans.extend(parse_inline(g["s_text"], Span(
                "", base.bold, base.italic, base.code, True, base.link)))
        elif g["br"]:
            spans.append(Span("", brk=True))
    plain(text[pos:])
    return spans


def plain_text(text: str) -> str:
    return "".join("\n" if s.brk else s.text for s in parse_inline(text))


# ── Markdown blocks ──────────────────────────────────────────────────


@dataclass
class Block:
    kind: str                 # heading | para | bullet | number | quote | code | table | image | hr | pagebreak | toc
    text: str = ""
    level: int = 0
    start: int = 1
    rows: list[list[str]] = field(default_factory=list)
    align: list[str] = field(default_factory=list)
    src: str = ""
    lines: list[str] = field(default_factory=list)


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_HR_RE = re.compile(r"^ {0,3}([-*_])( *\1){2,} *$")
_LIST_RE = re.compile(r"^(\s*)([-*+•]|\d{1,3}[.)])\s+(.*)$")
_IMAGE_RE = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<src>[^)\s]+)(?:\s+\"(?P<title>[^\"]*)\")?\)\s*$")
_FENCE_RE = re.compile(r"^\s*(```|~~~)")
_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_PAGEBREAK = {"\\pagebreak", "\\newpage", "<!-- pagebreak -->", "<div style=\"page-break-after: always\"></div>"}


def _table_cells(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", line)]


def parse_blocks(markdown: str) -> list[Block]:
    lines = (markdown or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[Block] = []
    i, n = 0, len(lines)
    indents: list[int] = []

    def table_starts(line: str, nxt: str) -> bool:
        return "|" in line and "|" in (nxt or "") and bool(_TABLE_SEP_RE.match(nxt or ""))

    def starts_block(line: str, nxt: str) -> bool:
        s = line.strip()
        return bool(
            _HEADING_RE.match(s) or _HR_RE.match(line) or _LIST_RE.match(line)
            or _FENCE_RE.match(line) or s.startswith(">") or _IMAGE_RE.match(s)
            or s in _PAGEBREAK or s.upper() == "[TOC]" or table_starts(line, nxt))

    while i < n:
        line = lines[i]
        stripped = line.strip()
        nxt = lines[i + 1] if i + 1 < n else ""
        if not stripped:
            i += 1
            continue
        list_m = _LIST_RE.match(line)
        if not list_m:
            indents = []
        if _FENCE_RE.match(line):
            fence = _FENCE_RE.match(line).group(1)
            body: list[str] = []
            i += 1
            while i < n and not lines[i].strip().startswith(fence):
                body.append(lines[i])
                i += 1
            blocks.append(Block("code", lines=body))
            i += 1
            continue
        if stripped in _PAGEBREAK:
            blocks.append(Block("pagebreak"))
            i += 1
            continue
        if stripped.upper() == "[TOC]":
            blocks.append(Block("toc"))
            i += 1
            continue
        m = _HEADING_RE.match(stripped)
        if m:
            blocks.append(Block("heading", m.group(2), level=len(m.group(1))))
            i += 1
            continue
        if _HR_RE.match(line):
            blocks.append(Block("hr"))
            i += 1
            continue
        m = _IMAGE_RE.match(stripped)
        if m:
            blocks.append(Block("image", m.group("title") or m.group("alt"), src=m.group("src")))
            i += 1
            continue
        if table_starts(line, nxt):
            header = _table_cells(line)
            align = []
            for cell in _table_cells(nxt):
                c = cell.strip()
                align.append("center" if c.startswith(":") and c.endswith(":")
                             else "right" if c.endswith(":") else "left")
            rows = [header]
            i += 2
            while i < n and "|" in lines[i] and lines[i].strip():
                rows.append(_table_cells(lines[i]))
                i += 1
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            blocks.append(Block("table", rows=rows, align=(align + ["left"] * width)[:width]))
            continue
        if stripped.startswith(">"):
            quote: list[str] = []
            while i < n and lines[i].strip().startswith(">"):
                quote.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            for chunk in "\n".join(quote).split("\n\n"):
                if chunk.strip():
                    blocks.append(Block("quote", _join_lines(chunk.split("\n"))))
            continue
        if list_m:
            indent = len(list_m.group(1).replace("\t", "    "))
            while indents and indents[-1] > indent:
                indents.pop()
            if not indents or indents[-1] < indent:
                indents.append(indent)
            level = min(len(indents) - 1, 2)
            marker = list_m.group(2)
            item = [list_m.group(3)]
            i += 1
            while i < n and lines[i].strip() and not starts_block(
                    lines[i], lines[i + 1] if i + 1 < n else ""):
                item.append(lines[i].strip())
                i += 1
            if marker[0].isdigit():
                blocks.append(Block("number", _join_lines(item), level=level,
                                    start=int(marker[:-1])))
            else:
                blocks.append(Block("bullet", _join_lines(item), level=level))
            continue
        para = [line]
        i += 1
        while i < n and lines[i].strip() and not starts_block(
                lines[i], lines[i + 1] if i + 1 < n else ""):
            para.append(lines[i])
            i += 1
        blocks.append(Block("para", _join_lines(para)))
    return blocks


def _join_lines(lines: list[str]) -> str:
    out = ""
    for i, line in enumerate(lines):
        hard = line.endswith("  ") or line.endswith("\\")
        piece = line.rstrip("\\").strip()
        out += piece
        if i < len(lines) - 1:
            out += "<br>" if hard else " "
    return out


# ── Building content ─────────────────────────────────────────────────


@dataclass
class Builder:
    """Turns Markdown blocks into paragraphs and tables of *doc*.

    *base_rpr* is run formatting new text starts from (an edit carries the
    surrounding text's font); *plain_ppr* the paragraph properties a plain
    paragraph takes (an edit inherits the paragraph it follows).
    """

    doc: DocxDocument
    image_root: Callable[[str], Path | None] | None = None
    base_rpr: Any | None = None
    plain_ppr: Any | None = None
    notes: list[str] = field(default_factory=list)
    _nums: dict[int, str] = field(default_factory=dict)
    _rpr: Any | None = None  # base_rpr, or only its fonts for a heading

    def style(self, name: str) -> Any | None:
        try:
            return self.doc.styles[name]
        except KeyError:
            return None

    def build(self, blocks: list[Block]) -> list[Any]:
        out: list[Any] = []
        for block in blocks:
            if block.kind not in ("number", "bullet"):
                self._nums.clear()
            self._rpr = self.base_rpr
            if block.kind in ("heading", "code") and self.base_rpr is not None:
                # A heading keeps its style's size; only the typeface carries.
                fonts = _el("w:rPr")
                for ch in self.base_rpr:
                    if etree.QName(ch).localname in ("rFonts", "lang") and block.kind == "heading":
                        fonts.append(copy.deepcopy(ch))
                self._rpr = fonts if len(fonts) else None
            out.extend(self._block(block))
        return out

    # paragraphs

    def _new_p(self, style: str | None = None) -> Paragraph:
        p = Paragraph(_el("w:p"), self.doc._body)
        if style and self.style(style) is not None:
            p.style = self.doc.styles[style]
        return p

    def _plain_p(self) -> Paragraph:
        p = Paragraph(_el("w:p"), self.doc._body)
        if self.plain_ppr is not None:
            ppr = copy.deepcopy(self.plain_ppr)
            for tag in (RPR, SECT_PR, PPR_CHANGE):
                for el in ppr.findall(tag):
                    ppr.remove(el)
            p._p.insert(0, ppr)
        return p

    def _runs(self, p: Paragraph, text: str, *, code: bool = False) -> None:
        for span in parse_inline(text):
            if span.brk:
                p.add_run().add_break()
                continue
            if span.link:
                self._link(p, span)
                continue
            run = p.add_run()
            self._dress(run, span, code=code)
            run.text = span.text

    def _dress(self, run: Run, span: Span, *, code: bool = False) -> None:
        if self._rpr is not None:
            run._r.insert(0, copy.deepcopy(self._rpr))
        font = run.font
        if span.bold:
            font.bold = True
        if span.italic:
            font.italic = True
        if span.strike:
            font.strike = True
        if span.code or code:
            font.name = "Consolas"
            rfonts = run._r.get_or_add_rPr().get_or_add_rFonts()
            rfonts.set(qn("w:hAnsi"), "Consolas")
            rfonts.set(qn("w:cs"), "Consolas")

    def _link(self, p: Paragraph, span: Span) -> None:
        url = span.link
        if not re.match(r"^(https?://|mailto:)", url, re.I):
            run = p.add_run()
            self._dress(run, span)
            run.text = span.text
            return
        r_id = self.doc.part.relate_to(url, RT.HYPERLINK, is_external=True)
        link = _el("w:hyperlink")
        link.set(qn("r:id"), r_id)
        run = Run(_el("w:r"), p)
        self._dress(run, span)
        if self.style("Hyperlink") is not None:
            run.style = self.doc.styles["Hyperlink"]
        else:
            run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
            run.font.underline = True
        run.text = span.text
        link.append(run._r)
        p._p.append(link)

    def _block(self, block: Block) -> list[Any]:
        kind = block.kind
        if kind == "heading":
            name = f"Heading {block.level}"
            if self.style(name) is None:
                p = self._new_p()
                self._runs(p, f"**{block.text}**")
            else:
                p = self._new_p(name)
                self._runs(p, block.text)
            return [p._p]
        if kind == "para":
            p = self._plain_p()
            self._runs(p, block.text)
            return [p._p]
        if kind in ("bullet", "number"):
            return [self._list_item(block)]
        if kind == "quote":
            style = next((s for s in ("Quote", "Intense Quote") if self.style(s)), None)
            p = self._new_p(style)
            if style is None:
                p.paragraph_format.left_indent = Inches(0.5)
                self._runs(p, f"*{block.text}*")
            else:
                self._runs(p, block.text)
            return [p._p]
        if kind == "code":
            p = self._new_p("No Spacing" if self.style("No Spacing") else None)
            shd = _el("w:shd")
            shd.set(W_VAL, "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "F2F2F2")
            ppr = p._p.get_or_add_pPr()
            ppr.append(shd)
            for idx, line in enumerate(block.lines):
                if idx:
                    p.add_run().add_break()
                run = p.add_run()
                self._dress(run, Span(""), code=True)
                run.font.size = Pt(9.5)
                run.text = line
            return [p._p]
        if kind == "table":
            return [self._table(block)]
        if kind == "image":
            return self._image(block)
        if kind == "hr":
            p = self._new_p()
            pbdr = _el("w:pBdr")
            bottom = _el("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "auto")):
                bottom.set(qn(k), v)
            pbdr.append(bottom)
            p._p.get_or_add_pPr().append(pbdr)
            return [p._p]
        if kind == "pagebreak":
            p = self._new_p()
            p.add_run().add_break(WD_BREAK.PAGE)
            return [p._p]
        if kind == "toc":
            return self._toc()
        return []

    def _list_item(self, block: Block) -> Any:
        numbered = block.kind == "number"
        base = "List Number" if numbered else "List Bullet"
        name = base if block.level == 0 else f"{base} {block.level + 1}"
        if self.style(name) is None:
            name = base
        if self.style(name) is None:
            p = self._plain_p()
            marker = f"{block.start}. " if numbered else "• "
            self._runs(p, marker + block.text)
            return p._p
        p = self._new_p(name)
        for lvl in list(self._nums):
            if lvl > block.level:
                del self._nums[lvl]
        if numbered:
            num_id = self._nums.get(block.level)
            if num_id is None:
                num_id = self._restart(name, block.start)
                if num_id:
                    self._nums[block.level] = num_id
            if num_id:
                numpr = p._p.get_or_add_pPr().get_or_add_numPr()
                numpr.get_or_add_ilvl().val = self._style_ilvl(name)
                numpr.get_or_add_numId().val = int(num_id)
        else:
            self._nums.pop(block.level, None)
        self._runs(p, block.text)
        return p._p

    def _style_numpr(self, name: str) -> Any | None:
        st = self.style(name)
        ppr = st.element.pPr if st is not None else None
        return ppr.numPr if ppr is not None else None

    def _style_ilvl(self, name: str) -> int:
        numpr = self._style_numpr(name)
        ilvl = numpr.ilvl if numpr is not None else None
        return int(ilvl.val) if ilvl is not None else 0

    def _restart(self, name: str, start: int) -> str:
        """A fresh w:num on the style's list definition, counting from *start*."""
        numpr = self._style_numpr(name)
        if numpr is None or numpr.numId is None:
            return ""
        try:
            numbering = self.doc.part.numbering_part.element
            num = numbering.num_having_numId(numpr.numId.val)
            new = numbering.add_num(num.abstractNumId.val)
            override = new.add_lvlOverride(ilvl=self._style_ilvl(name))
            override.add_startOverride(max(1, int(start)))
            return str(new.numId)
        except Exception:
            return ""

    def _table(self, block: Block) -> Any:
        from docx.oxml.table import CT_Tbl

        rows, cols = len(block.rows), len(block.rows[0])
        width = self.doc._block_width
        tbl = CT_Tbl.new_tbl(rows, cols, width)
        table = Table(tbl, self.doc._body)
        for name in ("Table Grid", "Light Grid Accent 1"):
            if self.style(name) is not None:
                table.style = self.doc.styles[name]
                break
        aligns = {"center": 1, "right": 2}
        for ri, row in enumerate(block.rows):
            for ci, text in enumerate(row):
                para = table.cell(ri, ci).paragraphs[0]
                self._runs(para, f"**{text}**" if ri == 0 and text.strip() else text)
                if block.align[ci] in aligns:
                    from docx.enum.text import WD_ALIGN_PARAGRAPH
                    para.alignment = (WD_ALIGN_PARAGRAPH.CENTER if block.align[ci] == "center"
                                      else WD_ALIGN_PARAGRAPH.RIGHT)
        # Repeat the header row on every page.
        tr = tbl.findall(TR)[0]
        trpr = tr.get_or_add_trPr()
        header = _el("w:tblHeader")
        trpr.append(header)
        return tbl

    def _image(self, block: Block) -> list[Any]:
        src = block.src
        if re.match(r"^[a-z]+://", src, re.I):
            self.notes.append(
                f"Picture {src} skipped: only saved files can go in (web_fetch the "
                "picture first; it is saved and listed with a path).")
            return []
        path = self.image_root(src) if self.image_root else Path(src)
        if path is None or not path.is_file():
            self.notes.append(f"Picture {src} skipped: file not found.")
            return []
        if path.suffix.lower() not in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff"}:
            self.notes.append(f"Picture {src} skipped: {path.suffix or 'no'} is not a picture format Word takes.")
            return []
        p = self._new_p()
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        try:
            shape = run.add_picture(str(path))
            limit = int(self.doc._block_width)
            if shape.width and shape.width > limit:
                ratio = limit / shape.width
                shape.width = Emu(limit)
                shape.height = Emu(int(shape.height * ratio))
        except Exception as exc:
            self.notes.append(f"Picture {src} skipped: {exc}.")
            return []
        out = [p._p]
        if block.text.strip():
            cap = self._new_p("Caption" if self.style("Caption") else None)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._runs(cap, block.text if self.style("Caption") else f"*{block.text}*")
            out.append(cap._p)
        return out

    def _toc(self) -> list[Any]:
        heading = self._new_p("TOC Heading" if self.style("TOC Heading") else None)
        heading.add_run("Contents")
        p = self._new_p()
        for kind, text in (("begin", None), ("instr", ' TOC \\o "1-3" \\h \\z \\u '),
                           ("separate", None), ("text", "Right-click and choose Update Field to build the table of contents."),
                           ("end", None)):
            r = _el("w:r")
            if kind == "instr":
                r.append(_text_el(text, "w:instrText"))
            elif kind == "text":
                r.append(_text_el(text))
            else:
                fc = _el("w:fldChar")
                fc.set(qn("w:fldCharType"), kind)
                if kind == "begin":
                    fc.set(qn("w:dirty"), "true")
                r.append(fc)
            p._p.append(r)
        _settings_flag(self.doc, "updateFields", True)
        return [heading._p, p._p]


# ── Settings ─────────────────────────────────────────────────────────

# w:settings children in schema order (ECMA-376 17.15.1.78).
_SETTINGS_ORDER = (
    "writeProtection view zoom removePersonalInformation removeDateAndTime "
    "doNotDisplayPageBoundaries displayBackgroundShape printPostScriptOverText "
    "printFractionalCharacterWidth printFormsData embedTrueTypeFonts embedSystemFonts "
    "saveSubsetFonts saveFormsData mirrorMargins alignBordersAndEdges "
    "bordersDoNotSurroundHeader bordersDoNotSurroundFooter gutterAtTop hideSpellingErrors "
    "hideGrammaticalErrors activeWritingStyle proofState formsDesign attachedTemplate "
    "linkStyles stylePaneFormatFilter stylePaneSortMethod documentType mailMerge "
    "revisionView trackRevisions doNotTrackMoves doNotTrackFormatting documentProtection "
    "autoFormatOverride styleLockTheme styleLockQFSet defaultTabStop autoHyphenation "
    "consecutiveHyphenLimit hyphenationZone doNotHyphenateCaps showEnvelope summaryLength "
    "clickAndTypeStyle defaultTableStyle evenAndOddHeaders bookFoldRevPrinting "
    "bookFoldPrinting bookFoldPrintingSheets drawingGridHorizontalSpacing "
    "drawingGridVerticalSpacing displayHorizontalDrawingGridEvery "
    "displayVerticalDrawingGridEvery doNotUseMarginsForDrawingGridOrigin "
    "drawingGridHorizontalOrigin drawingGridVerticalOrigin doNotShadeFormData "
    "noPunctuationKerning characterSpacingControl printTwoOnOne strictFirstAndLastChars "
    "noLineBreaksAfter noLineBreaksBefore savePreviewPicture doNotValidateAgainstSchema "
    "saveInvalidXml ignoreMixedContent alwaysShowPlaceholderText doNotDemarcateInvalidXml "
    "saveXmlDataOnly useXSLTWhenSaving saveThroughXslt showXMLTags "
    "alwaysMergeEmptyNamespace updateFields hdrShapeDefaults footnotePr endnotePr compat "
    "docVars rsids"
).split()


def _settings_flag(doc: DocxDocument, name: str, on: bool) -> None:
    settings = doc.settings.element
    tag = qn(f"w:{name}")
    existing = settings.find(tag)
    if not on:
        if existing is not None:
            settings.remove(existing)
        return
    if existing is not None:
        existing.attrib.pop(W_VAL, None)
        return
    el = _el(f"w:{name}")
    later = set(_SETTINGS_ORDER[_SETTINGS_ORDER.index(name) + 1:])
    for child in settings:
        local = etree.QName(child).localname
        if etree.QName(child).namespace != W_NS or local in later:
            child.addprevious(el)
            return
    settings.append(el)


def tracking_on(doc: DocxDocument) -> bool:
    el = doc.settings.element.find(qn("w:trackRevisions"))
    return el is not None and el.get(W_VAL, "true") not in ("false", "0", "off")


def set_tracking(doc: DocxDocument, on: bool) -> None:
    _settings_flag(doc, "trackRevisions", on)


# ── Creating a document ──────────────────────────────────────────────


@dataclass
class Look:
    font: str = ""
    font_size: float = 0
    heading_font: str = ""
    heading_color: str = ""


def _style_font(style: Any, name: str = "", size: float = 0, color: str = "") -> None:
    if name:
        style.font.name = name
        rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), name)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
    if size:
        style.font.size = Pt(size)
    if color:
        style.font.color.rgb = RGBColor.from_string(color.lstrip("#").upper())


def build_document(
    markdown: str,
    *,
    template: str | Path | None = None,
    title: str = "",
    author: str = "",
    look: Look | None = None,
    page_numbers: bool = False,
    track_changes: bool = False,
    image_root: Callable[[str], Path | None] | None = None,
) -> tuple[DocxDocument, list[str]]:
    """A new document from Markdown (and optionally a template's styles,
    page setup, headers and footers). Returns it with notes for the model."""
    import docx as _docx

    if template:
        doc = open_document(template, template=True)
        body = doc.element.body
        for child in list(body):
            if child.tag != SECT_PR:
                body.remove(child)
    else:
        doc = _docx.Document()
    look = look or Look()
    styles = doc.styles
    if look.font or look.font_size:
        try:
            _style_font(styles["Normal"], look.font, look.font_size)
        except KeyError:
            pass
    if look.heading_font or look.heading_color:
        for name in ("Title", *(f"Heading {i}" for i in range(1, 7))):
            try:
                _style_font(styles[name], look.heading_font, 0, look.heading_color)
            except KeyError:
                continue
    builder = Builder(doc, image_root=image_root)
    blocks = parse_blocks(markdown)
    if not blocks:
        raise DocxError("content is empty — pass the document text as Markdown.")
    body = doc.element.body
    sect = body.find(SECT_PR)
    for el in builder.build(blocks):
        if sect is not None:
            sect.addprevious(el)
        else:
            body.append(el)
    if title:
        doc.core_properties.title = title
    if author:
        doc.core_properties.author = author
        doc.core_properties.last_modified_by = author
    if page_numbers:
        _add_page_numbers(doc)
    if track_changes:
        set_tracking(doc, True)
    return doc, builder.notes


def _add_page_numbers(doc: DocxDocument) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        if visible_text(p._p).strip():
            continue
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for kind, value in (("text", "Page "), ("field", "PAGE"), ("text", " of "), ("field", "NUMPAGES")):
            if kind == "text":
                p.add_run(value)
                continue
            fld = _el("w:fldSimple")
            fld.set(qn("w:instr"), f" {value} ")
            r = _el("w:r")
            r.append(_text_el("1"))
            fld.append(r)
            p._p.append(fld)


# ── Edits ────────────────────────────────────────────────────────────

OPS = ("replace", "delete", "insert", "insert_paragraphs", "format", "set_style", "comment")
_FORMAT_KEYS = ("bold", "italic", "underline", "strike", "highlight", "color", "size")


@dataclass
class EditReport:
    lines: list[str] = field(default_factory=list)
    touched: list[Any] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


def _pick(hits: list[Hit], op: dict[str, Any], paras: list[Any], label: str) -> list[Hit]:
    needle = str(op.get("find") or "")
    if not hits:
        raise DocxError(_not_found(paras, needle))
    if op.get("all"):
        return hits
    occurrence = op.get("occurrence")
    if occurrence:
        try:
            k = int(occurrence)
        except (TypeError, ValueError):
            raise DocxError("occurrence must be a number (1 = first).") from None
        if not 1 <= k <= len(hits):
            raise DocxError(f"occurrence {k} is out of range: {needle[:60]!r} occurs {len(hits)} time(s).")
        return [hits[k - 1]]
    if len(hits) > 1:
        per_para: dict[int, int] = {}
        for h in hits:
            n = paras.index(h.p) + 1
            per_para[n] = per_para.get(n, 0) + 1
        where = ", ".join(f"[{n}]" + (f" ×{c}" if c > 1 else "")
                          for n, c in list(per_para.items())[:10])
        raise DocxError(
            f"{needle[:80]!r} occurs {len(hits)} times (paragraphs {where}). "
            "Add surrounding words to make it unique, or pass paragraph=N, "
            "occurrence=N, or all=true.")
    return hits


def _scope(op: dict[str, Any], snapshot: list[Any]) -> list[Any] | None:
    if op.get("paragraph") in (None, ""):
        return None
    return [_para(snapshot, op.get("paragraph"))]


def _para(snapshot: list[Any], number: Any) -> Any:
    try:
        n = int(number)
    except (TypeError, ValueError):
        raise DocxError(f"paragraph must be a paragraph number from `read`, got {number!r}.") from None
    if not 1 <= n <= len(snapshot):
        raise DocxError(f"paragraph {n} doesn't exist — the document has {len(snapshot)} paragraphs.")
    return snapshot[n - 1]


def _base_rpr(p: Any, *, keep: tuple[str, ...]) -> Any | None:
    for r in _own_runs(p):
        if _run_state(r, p) == "del" or not _run_text(r).strip():
            continue
        rpr = r.find(RPR)
        if rpr is None:
            return None
        base = _el("w:rPr")
        for ch in rpr:
            if etree.QName(ch).localname in keep:
                base.append(copy.deepcopy(ch))
        return base if len(base) else None
    return None


def _is_heading(doc: DocxDocument, p: Any) -> tuple[bool, Any | None]:
    para = Paragraph(p, doc._body)
    try:
        style = para.style
    except Exception:
        return False, None
    name = (style.name or "") if style is not None else ""
    return (name.startswith("Heading") or name in ("Title", "Subtitle")), style


def apply_edits(
    doc: DocxDocument,
    edits: list[dict[str, Any]],
    rv: Reviser,
    *,
    image_root: Callable[[str], Path | None] | None = None,
) -> EditReport:
    """Apply *edits* in order as tracked changes. Paragraph numbers refer to
    the document as it was before this call; `find` sees the text as it reads
    after the edits before it. Raises DocxError (nothing is saved then)."""
    snapshot = paragraphs(doc)
    report = EditReport()
    for index, op in enumerate(edits, 1):
        if not isinstance(op, dict):
            raise DocxError(f"edit {index} must be an object with an `op`.")
        kind = str(op.get("op") or "").strip().lower()
        label = f"edit {index} ({kind or '?'})"
        try:
            if kind == "replace":
                _op_replace(doc, op, snapshot, rv, report, label)
            elif kind == "delete":
                _op_delete(doc, op, snapshot, rv, report, label)
            elif kind == "insert":
                _op_insert(doc, op, snapshot, rv, report, label)
            elif kind == "insert_paragraphs":
                _op_insert_paragraphs(doc, op, snapshot, rv, report, image_root)
            elif kind == "format":
                _op_format(doc, op, snapshot, rv, report, label)
            elif kind == "set_style":
                _op_set_style(doc, op, snapshot, rv, report)
            elif kind == "comment":
                _op_comment(doc, op, snapshot, rv, report, label)
            else:
                raise DocxError(f"unknown op {kind!r}; use one of: {', '.join(OPS)}.")
        except DocxError as exc:
            raise DocxError(f"{label}: {exc}") from None
    return report


def _hits_for(doc: DocxDocument, op: dict[str, Any], snapshot: list[Any], label: str) -> list[Hit]:
    needle = op.get("find")
    if not isinstance(needle, str) or not needle:
        raise DocxError("`find` (the exact text to change) is required.")
    live = paragraphs(doc)
    hits = find_text(live, needle, _scope(op, snapshot))
    return _pick(hits, op, live, label)


def _op_replace(doc, op, snapshot, rv, report, label) -> None:
    if "replace_with" not in op or not isinstance(op.get("replace_with"), str):
        raise DocxError("`replace_with` is required (use op=delete to remove text).")
    new = op["replace_with"]
    for hit in reversed(_hits_for(doc, op, snapshot, label)):
        if visible_text(hit.p)[hit.start:hit.end] == new:
            continue
        _track_replace(hit.p, hit.start, hit.end, new, rv)
        report.touched.append(hit.p)


def _op_delete(doc, op, snapshot, rv, report, label) -> None:
    if op.get("find"):
        for hit in reversed(_hits_for(doc, op, snapshot, label)):
            _track_delete(_isolate(hit.p, hit.start, hit.end), rv)
            report.touched.append(hit.p)
        return
    if op.get("paragraph") in (None, ""):
        raise DocxError("give `find` (text to delete) or `paragraph` (with optional `to`) to delete whole paragraphs.")
    first = int(op.get("paragraph"))
    last = int(op.get("to") or first)
    if last < first:
        raise DocxError("`to` must be at or after `paragraph`.")
    for n in range(first, last + 1):
        p = _para(snapshot, n)
        if p.getparent() is None:
            continue
        runs = [r for r, _, _ in _segments(p)]
        if runs:
            _track_delete(runs, rv)
        state, marker = _mark_state(p)
        if state == "ins" and marker is not None and marker.get(W_AUTHOR) == rv.author \
                and not _own_runs(p):
            # A paragraph this author inserted: deleting it leaves no trace.
            p.getparent().remove(p)
            rv.counts["paragraphs_deleted"] += 1
            continue
        # The last paragraph of a cell or of the document keeps its mark.
        if any(s.tag in (P, TBL) for s in p.itersiblings()):
            _mark_paragraph(p, "del", rv)
        rv.counts["paragraphs_deleted"] += 1
        report.touched.append(p)


def _op_insert(doc, op, snapshot, rv, report, label) -> None:
    text = op.get("text")
    if not isinstance(text, str) or not text:
        raise DocxError("`text` (what to insert) is required.")
    where = "after" if op.get("after") else "before" if op.get("before") else ""
    if not where:
        raise DocxError("give `after` or `before`: the existing text the insertion goes next to.")
    anchor = dict(op, find=op.get(where))
    for hit in reversed(_hits_for(doc, anchor, snapshot, label)):
        pos = hit.end if where == "after" else hit.start
        marker, ref = _marker_at(hit.p, pos)
        _insert_at_marker(marker, text, _rpr_of(ref), rv)
        report.touched.append(hit.p)


def _op_insert_paragraphs(doc, op, snapshot, rv, report, image_root) -> None:
    content = op.get("content") if isinstance(op.get("content"), str) else op.get("text")
    if not isinstance(content, str) or not content.strip():
        raise DocxError("`content` (Markdown for the new paragraphs) is required.")
    if op.get("after") not in (None, ""):
        anchor, where = _para(snapshot, op.get("after")), "after"
    elif op.get("before") not in (None, ""):
        anchor, where = _para(snapshot, op.get("before")), "before"
    else:
        raise DocxError("give `after` or `before`: the paragraph number the new paragraphs go next to.")
    heading, style = _is_heading(doc, anchor)
    plain_ppr = None
    if not heading:
        plain_ppr = anchor.find(PPR)
    keep = ("rFonts", "lang") if heading else ("rFonts", "sz", "szCs", "lang")
    builder = Builder(doc, image_root=image_root, plain_ppr=plain_ppr,
                      base_rpr=_base_rpr(anchor, keep=keep))
    if heading and style is not None:
        nxt = style.next_paragraph_style
        if nxt is not None and nxt.style_id != style.style_id:
            ppr = _el("w:pPr")
            pstyle = _el("w:pStyle")
            pstyle.set(W_VAL, nxt.style_id)
            ppr.append(pstyle)
            builder.plain_ppr = ppr
    if op.get("style"):
        st = _style_named(doc, str(op.get("style")))
        ppr = _el("w:pPr")
        pstyle = _el("w:pStyle")
        pstyle.set(W_VAL, st.style_id)
        ppr.append(pstyle)
        builder.plain_ppr = ppr
    blocks = builder.build(parse_blocks(content))
    report.notes.extend(builder.notes)
    if not blocks:
        raise DocxError("nothing to insert.")
    cursor = anchor
    for block in blocks:
        _track_new_block(block, rv)
        if where == "after":
            cursor.addnext(block)
            cursor = block
        else:
            anchor.addprevious(block)
        rv.counts["paragraphs_inserted"] += 1 if block.tag == P else 0
        report.touched.append(block)
    # A table can't end its container: Word needs a paragraph after it.
    last = blocks[-1]
    if last.tag == TBL and (last.getnext() is None or last.getnext().tag == SECT_PR):
        tail = _el("w:p")
        last.addnext(tail)
        _mark_paragraph(tail, "ins", rv)


def _op_format(doc, op, snapshot, rv, report, label) -> None:
    props = {k: op[k] for k in _FORMAT_KEYS if k in op}
    if not props:
        raise DocxError(f"say what to change: any of {', '.join(_FORMAT_KEYS)}.")
    if op.get("find"):
        for hit in reversed(_hits_for(doc, op, snapshot, label)):
            _track_format(_isolate(hit.p, hit.start, hit.end), props, rv)
            report.touched.append(hit.p)
        return
    p = _para(snapshot, op.get("paragraph"))
    _track_format([r for r, s, e in _segments(p) if e > s], props, rv)
    report.touched.append(p)


def _style_named(doc: DocxDocument, name: str) -> Any:
    from docx.enum.style import WD_STYLE_TYPE

    for style in doc.styles:
        if style.type == WD_STYLE_TYPE.PARAGRAPH and (
                style.name == name or (style.name or "").lower() == name.lower()
                or style.style_id == name):
            return style
    names = sorted(s.name for s in doc.styles if s.type == WD_STYLE_TYPE.PARAGRAPH and s.name)
    raise DocxError(f"no paragraph style {name!r} in this document. Styles: {', '.join(names[:40])}")


def _op_set_style(doc, op, snapshot, rv, report) -> None:
    style = _style_named(doc, str(op.get("style") or ""))
    first = _para(snapshot, op.get("paragraph"))
    last = _para(snapshot, op.get("to") or op.get("paragraph"))
    i, j = snapshot.index(first), snapshot.index(last)
    for p in snapshot[i:j + 1]:
        _track_style(p, style, rv)
        report.touched.append(p)


def _op_comment(doc, op, snapshot, rv, report, label) -> None:
    text = op.get("text") or op.get("comment")
    if not isinstance(text, str) or not text.strip():
        raise DocxError("`text` (the comment) is required.")
    if op.get("find"):
        hits = _hits_for(doc, op, snapshot, label)
        targets = [(h.p, _isolate(h.p, h.start, h.end)) for h in hits]
    else:
        p = _para(snapshot, op.get("paragraph"))
        targets = [(p, [r for r, s, e in _segments(p) if e > s])]
    for p, runs in targets:
        if not runs:
            raise DocxError("there is no text there to anchor a comment on.")
        para = Paragraph(p, doc._body)
        doc.add_comment([Run(runs[0], para), Run(runs[-1], para)], text=text.strip(),
                        author=rv.author, initials=initials_of(rv.author))
        rv.counts["comments"] += 1
        report.touched.append(p)


# ── Accept / reject ──────────────────────────────────────────────────


def _is_mark(el: Any) -> bool:
    parent = el.getparent()
    return parent is not None and parent.tag == RPR and parent.getparent() is not None \
        and parent.getparent().tag == PPR


def _is_row_mark(el: Any) -> bool:
    parent = el.getparent()
    return parent is not None and parent.tag == TRPR


def _unwrap(el: Any) -> None:
    for ch in list(el):
        el.addprevious(ch)
    el.getparent().remove(el)


def resolve(doc: DocxDocument, *, accept: bool, ids: Iterable[str] | None = None,
            author: str = "") -> dict[str, int]:
    """Accept or reject revisions — all, those with *ids*, or *author*'s."""
    root = doc.element
    body = root.body
    wanted = {str(i).strip().lstrip("#") for i in ids} if ids else None
    who = (author or "").strip().lower()

    def chosen(el: Any) -> bool:
        if not _attached(el, root):
            return False
        if wanted is not None:
            return el.get(W_ID) in wanted
        if who:
            return (el.get(W_AUTHOR) or "").strip().lower() == who
        return True

    done = {"accepted" if accept else "rejected": 0}
    key = next(iter(done))

    # Text: insertions and deletions (paragraph-mark and row markers later).
    for el in list(body.iter(INS, DEL, MOVE_FROM, MOVE_TO)):
        if _is_mark(el) or _is_row_mark(el) or not chosen(el):
            continue
        keep = (el.tag in SHOWN_CONTAINERS) == accept
        if keep:
            if el.tag in HIDDEN_CONTAINERS:
                _to_live(el)
            _unwrap(el)
        else:
            el.getparent().remove(el)
        done[key] += 1

    # Property changes.
    for el in list(body.iter(*PROPERTY_CHANGES)):
        if not chosen(el):
            continue
        parent = el.getparent()
        parent.remove(el)
        if not accept and len(el):
            lead, tail = PROPERTY_CHANGES[el.tag]
            for ch in list(parent):
                if ch.tag not in lead and ch.tag not in tail:
                    parent.remove(ch)
            position = sum(1 for ch in parent if ch.tag in lead)
            for offset, ch in enumerate(list(el[0])):
                parent.insert(position + offset, ch)
        done[key] += 1

    # Paragraph marks, in document order (a run of inserted paragraphs folds
    # into the paragraph after it).
    for el in list(body.iter(INS, DEL, MOVE_FROM, MOVE_TO)):
        if not _is_mark(el) or not chosen(el):
            continue
        p = el.getparent().getparent().getparent()
        rpr = el.getparent()
        rpr.remove(el)
        if not len(rpr):
            rpr.getparent().remove(rpr)
        if (el.tag in HIDDEN_CONTAINERS) == accept:
            _drop_mark(p)
        done[key] += 1

    # Table rows and cells.
    for el in list(body.iter(INS, DEL, CELL_INS, CELL_DEL)):
        if not chosen(el):
            continue
        parent = el.getparent()
        if el.tag in (INS, DEL) and parent.tag == TRPR:
            tr = parent.getparent()
            parent.remove(el)
            if (el.tag == DEL) == accept:
                tbl = tr.getparent()
                tbl.remove(tr)
                if not tbl.findall(TR):
                    tbl.getparent().remove(tbl)
            done[key] += 1
        elif el.tag in (CELL_INS, CELL_DEL) and parent.tag == TCPR:
            tc = parent.getparent()
            parent.remove(el)
            if (el.tag == CELL_DEL) == accept:
                tc.getparent().remove(tc)
            done[key] += 1

    if not any(True for _ in body.iter(MOVE_FROM, MOVE_TO)):
        for tag in ("moveFromRangeStart", "moveFromRangeEnd", "moveToRangeStart", "moveToRangeEnd"):
            for el in list(body.iter(qn(f"w:{tag}"))):
                el.getparent().remove(el)
    _prune_comments(doc)
    return done


def _prune_comments(doc: DocxDocument) -> None:
    """Drop comment anchors whose reference went with removed text."""
    body = doc.element.body
    referenced = {el.get(W_ID) for el in body.iter(COMMENT_REF)}
    for el in list(body.iter(CRS, CRE)):
        if el.get(W_ID) not in referenced:
            el.getparent().remove(el)


def has_comments_part(doc: DocxDocument) -> bool:
    """Without creating one (python-docx's ``doc.comments`` adds the part)."""
    try:
        doc.part.part_related_by(RT.COMMENTS)
        return True
    except KeyError:
        return False


# ── Reading ──────────────────────────────────────────────────────────


@dataclass
class Change:
    id: str
    kind: str
    author: str
    date: str
    paragraph: int
    text: str


def list_changes(doc: DocxDocument) -> list[Change]:
    body = doc.element.body
    paras = paragraphs(doc)
    index = {p: i for i, p in enumerate(paras, 1)}

    def para_of(el: Any) -> int:
        node = el
        while node is not None and node is not body:
            if node.tag == P and node in index:
                return index[node]
            node = node.getparent()
        # A table-row marker: its first paragraph.
        first = next(iter(el.getparent().getparent().iter(P)), None) if el.getparent() is not None else None
        return index.get(first, 0)

    out: list[Change] = []
    names = {INS: "insertion", DEL: "deletion", MOVE_FROM: "moved away", MOVE_TO: "moved here"}
    for el in body.iter(INS, DEL, MOVE_FROM, MOVE_TO, CELL_INS, CELL_DEL, *PROPERTY_CHANGES):
        tag = el.tag
        if tag in (INS, DEL, MOVE_FROM, MOVE_TO):
            if _is_mark(el):
                kind, text = ("paragraph inserted" if tag in SHOWN_CONTAINERS
                              else "paragraph break deleted"), ""
            elif _is_row_mark(el):
                kind = "table row inserted" if tag == INS else "table row deleted"
                text = " | ".join(visible_text(p) for p in el.getparent().getparent().iter(P))
            else:
                kind = names[tag]
                text = "".join(_run_text(r) for r in el.iter(R))
        elif tag in (CELL_INS, CELL_DEL):
            kind, text = ("cell inserted" if tag == CELL_INS else "cell deleted"), ""
        elif tag == RPR_CHANGE:
            run = el.getparent().getparent()
            kind = "formatting"
            text = _run_text(run) if run.tag == R else ""
        elif tag == PPR_CHANGE:
            kind = "paragraph style"
            old_style = el.find(f"{PPR}/{qn('w:pStyle')}")
            new_style = el.getparent().find(qn("w:pStyle"))
            text = f"{old_style.get(W_VAL) if old_style is not None else 'Normal'} → " \
                   f"{new_style.get(W_VAL) if new_style is not None else 'Normal'}"
        else:
            kind, text = etree.QName(tag).localname, ""
        out.append(Change(el.get(W_ID) or "?", kind, el.get(W_AUTHOR) or "", el.get(W_DATE) or "",
                          para_of(el), text))
    return out


@dataclass
class CommentInfo:
    id: str
    author: str
    date: str
    text: str
    anchor: str
    paragraph: int


def list_comments(doc: DocxDocument) -> list[CommentInfo]:
    if not has_comments_part(doc):
        return []
    try:
        comments = list(doc.comments)
    except Exception:
        return []
    if not comments:
        return []
    body = doc.element.body
    paras = paragraphs(doc)
    index = {p: i for i, p in enumerate(paras, 1)}
    anchors: dict[str, list[str]] = {}
    where: dict[str, int] = {}
    open_ids: set[str] = set()
    for el in body.iter(CRS, CRE, T, P):
        if el.tag == P:
            continue
        if el.tag == CRS:
            cid = el.get(W_ID)
            open_ids.add(cid)
            anchors.setdefault(cid, [])
            node = el
            while node is not None and node.tag != P:
                node = node.getparent()
            where[cid] = index.get(node, 0)
        elif el.tag == CRE:
            open_ids.discard(el.get(W_ID))
        elif open_ids:
            for cid in open_ids:
                anchors[cid].append(el.text or "")
    out = []
    for c in comments:
        cid = str(c.comment_id)
        try:
            when = c.timestamp.strftime("%Y-%m-%d %H:%M") if c.timestamp else ""
        except Exception:
            when = ""
        text = "\n".join(p.text for p in c.paragraphs).strip()
        out.append(CommentInfo(cid, c.author or "", when, text,
                               "".join(anchors.get(cid, [])), where.get(cid, 0)))
    return out


def _para_label(p: Any, tables: dict[Any, int], style_names: dict[str, str]) -> str:
    bits = []
    tc = p.getparent()
    if tc is not None and tc.tag == TC:
        tr = tc.getparent()
        tbl = tr.getparent()
        rows = tbl.findall(TR)
        cells = tr.findall(TC)
        bits.append(f"T{tables.get(tbl, 0)} r{rows.index(tr) + 1}c{cells.index(tc) + 1}")
    ppr = p.find(PPR)
    pstyle = ppr.find(qn("w:pStyle")) if ppr is not None else None
    name = style_names.get(pstyle.get(W_VAL), pstyle.get(W_VAL)) if pstyle is not None else ""
    if name.startswith("Heading "):
        bits.append("H" + name.split(" ", 1)[1])
    elif name and name not in ("Normal", "Body Text"):
        bits.append(name)
    if ppr is not None and ppr.find(qn("w:numPr")) is not None and "List" not in name:
        bits.append("numbered")
    return f"({', '.join(bits)}) " if bits else ""


def labeler(doc: DocxDocument) -> Callable[[Any], str]:
    """``label(p)`` — the "(H2) " / "(T1 r2c3) " prefix `render` puts on *p*."""
    tables = {tbl: i for i, tbl in enumerate(doc.element.body.iter(TBL), 1)}
    names = {st.style_id: st.name or st.style_id for st in doc.styles}
    return lambda p: _para_label(p, tables, names)


def render_paragraph(p: Any, view: str = "markup", comment_ids: bool = True) -> str:
    out: list[str] = []
    state_open = ""
    for el in p.iter(R, CRS, CRE):
        if el.tag in (CRS, CRE):
            if view != "markup" or not comment_ids:
                continue
            node, own = el.getparent(), True
            while node is not None and node is not p:
                if node.tag in (P, TXBX):
                    own = False
                    break
                node = node.getparent()
            if not own:
                continue
            if state_open:
                out.append("+}" if state_open == "ins" else "-]")
                state_open = ""
            out.append("⟦" if el.tag == CRS else f"⟧c{el.get(W_ID)}")
            continue
        node, own = el.getparent(), True
        while node is not None and node is not p:
            if node.tag in (P, TXBX):
                own = False
                break
            node = node.getparent()
        if not own:
            continue
        state = _run_state(el, p)
        text = _run_text(el)
        if not text:
            continue
        if view == "final" and state == "del" or view == "original" and state == "ins":
            continue
        if view != "markup":
            out.append(text)
            continue
        if state != state_open:
            if state_open:
                out.append("+}" if state_open == "ins" else "-]")
            if state:
                out.append("{+" if state == "ins" else "[-")
            state_open = state
        out.append(text)
    if state_open:
        out.append("+}" if state_open == "ins" else "-]")
    text = "".join(out)
    if view == "markup":
        mark, _ = _mark_state(p)
        if mark == "ins":
            text += " ¶+"
        elif mark == "del":
            text += " ¶−"
        if any(True for _ in p.iter(RPR_CHANGE)) or p.find(f"{PPR}/{PPR_CHANGE}") is not None:
            text += " ‹format change›"
    return text


def summary(doc: DocxDocument) -> dict[str, Any]:
    changes = list_changes(doc)
    by_author: dict[str, int] = {}
    for c in changes:
        by_author[c.author or "?"] = by_author.get(c.author or "?", 0) + 1
    return {
        "paragraphs": len(paragraphs(doc)),
        "tables": len(list(doc.element.body.iter(TBL))),
        "sections": len(doc.sections),
        "tracking": tracking_on(doc),
        "changes": len(changes),
        "by_author": by_author,
        "comments": len(list_comments(doc)),
    }


def render(doc: DocxDocument, *, view: str = "markup", offset: int = 1, limit: int = 200,
           max_chars: int = 40000) -> tuple[str, int]:
    """Numbered paragraphs (empty ones left out, numbering kept). Returns the
    text and the next paragraph number to read from (0 when done)."""
    paras = paragraphs(doc)
    tables = {tbl: i for i, tbl in enumerate(doc.element.body.iter(TBL), 1)}
    style_names = {st.style_id: st.name or st.style_id for st in doc.styles}
    comments = {c.id: c for c in list_comments(doc)} if view == "markup" else {}
    lines: list[str] = []
    used = 0
    shown = 0
    start = max(1, int(offset or 1))
    for n in range(start, len(paras) + 1):
        p = paras[n - 1]
        text = render_paragraph(p, view)
        if not text.strip():
            continue
        line = f"[{n}] {_para_label(p, tables, style_names)}{text}"
        notes = []
        if comments:
            for cid in re.findall(r"⟧c(\w+)", text):
                c = comments.get(cid)
                if c:
                    notes.append(f"      comment c{cid} — {c.author}: {c.text}")
        block = "\n".join([line, *notes])
        if shown >= limit or (used + len(block) > max_chars and shown):
            return "\n".join(lines), n
        lines.append(block)
        used += len(block) + 1
        shown += 1
    return "\n".join(lines), 0


def headers_and_footers(doc: DocxDocument) -> list[str]:
    seen: list[str] = []
    for section in doc.sections:
        for label, part in (("Header", section.header), ("Footer", section.footer)):
            try:
                if part.is_linked_to_previous and seen:
                    continue
                text = " / ".join(t for t in (visible_text(p._p).strip() for p in part.paragraphs) if t)
            except Exception:
                continue
            if text and f"{label}: {text}" not in seen:
                seen.append(f"{label}: {text}")
    return seen
